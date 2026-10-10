
import streamlit as st
import pandas as pd
import numpy as np
import yfinance as yf
import FinanceDataReader as fdr
import concurrent.futures
import datetime

# 한국 표준시(KST) 타임존 (UTC+9)
KST = datetime.timezone(datetime.timedelta(hours=9))
import io
import plotly.graph_objects as go
from plotly.subplots import make_subplots

import time
import importlib
import sys
import os

# Streamlit Cloud 환경에서 로컬 모듈 캐시 갱신 보장 및 상세 오류 트래킹
try:
    import tickers
    importlib.reload(tickers)
    from tickers import get_krx_tickers

    import screener
    importlib.reload(screener)
    from screener import run_screening_task, run_screener, fit_upper_trendline, screen_single_stock
except Exception as e:
    import traceback
    st.error(f"모듈 로드 중 오류가 발생했습니다: {e}")
    st.code(traceback.format_exc())
    raise e

STANDARD_CHART_THEME = {
    'paper_bgcolor': '#1E293B',    # Tailwind Slate-800 (외곽 카드 배경)
    'plot_bgcolor': '#0F172A',     # Tailwind Slate-900 (내부 딥 블랙 플롯)
    'text_main': '#F8FAFC',        # 타이틀/헤더 텍스트 (순백색)
    'text_body': '#E2E8F0',        # 본문 및 축 라벨 (부드러운 화이트)
    'text_muted': '#CBD5E1',       # 축 눈금 수치 텍스트 (Slate-300)
    'grid_color': '#334155',       # 그리드 격자선 (Slate-700)
    'border_color': '#475569',     # 축 기준선 (Slate-600)
    'legend_bg': 'rgba(30, 41, 59, 0.85)',
    'legend_border': '#334155',
    'hover_bg': 'rgba(15, 23, 42, 0.9)',
    'hover_border': '#334155'
}

def fmt_curr(val, ticker):
    if ticker.endswith('.KS') or ticker.endswith('.KQ'):
        return f"{val:,.0f}원"
    else:
        return f"${val:,.2f}"

# 페이지 설정
st.set_page_config(
    page_title="David Ryan's Just Draw the Line Stock Screener",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# 커스텀 CSS로 UI 스타일링 (다크 테마 최적화 및 시인성 개선)
st.markdown("""
<style>
    /* Streamlit 고정 상단 헤더 배경 투명화 */
    header[data-testid="stHeader"] {
        background: transparent !important;
    }

    .main .block-container,
    [data-testid="stMainBlockContainer"],
    .block-container {
        padding-top: 2.0rem !important;
    }
    .main-title {
        font-size: 2.0rem !important;
        font-weight: 800 !important;
        color: #8AB4F8 !important;
        -webkit-text-fill-color: #8AB4F8 !important;
        margin-bottom: 0.2rem;
        text-align: center !important;
    }
    .sub-title {
        font-size: 0.92rem;
        color: #BDC1C6; /* 밝은 회색으로 가독성 향상 */
        margin-bottom: 2rem;
        text-align: center;
    }
    .metric-card {
        background-color: #202124; /* 검정색 계열의 배경 적용 */
        color: #CBD5E1; /* 눈 피로도 완화를 위한 부드러운 텍스트 색상 */
        padding: 15px;
        border-radius: 8px;
        border-left: 5px solid #8AB4F8; /* 하늘색 테두리 포인트 */
        margin-bottom: 10px;
    }
    .metric-card ul, .metric-card li {
        font-size: 0.9rem;
        line-height: 1.5;
    }
    /* 안내문(Alert) 박스 스타일: 폰트 및 이모지 아이콘 크기 축소 */
    .stAlert {
        padding: 0.5rem 0.85rem !important;
    }
    .stAlert [data-testid="stAlertDynamicIcon"],
    .stAlert [data-testid="stAlertDynamicIcon"] * {
        font-size: 1.05rem !important;
        width: 1.05rem !important;
        height: 1.05rem !important;
        line-height: 1 !important;
    }
    .stAlert svg {
        width: 1.05rem !important;
        height: 1.05rem !important;
    }
    .stAlert [data-testid="stMarkdownContainer"] p,
    .stAlert [data-testid="stMarkdownContainer"] span {
        font-size: 0.88rem !important;
    }
    /* 다운로드 버튼 공통 통일 스타일 */
    div[data-testid="stDownloadButton"] > button,
    .stDownloadButton > button {
        background-color: #334155 !important;
        color: #f8fafc !important;
        border: 1px solid #475569 !important;
        border-radius: 6px !important;
        font-size: 0.875rem !important;
        font-weight: 500 !important;
        height: 38px !important;
        min-height: 38px !important;
        max-height: 38px !important;
        line-height: 36px !important;
        padding: 0 16px !important;
        display: inline-flex !important;
        align-items: center !important;
        justify-content: center !important;
        text-align: center !important;
        transition: all 0.2s ease-in-out !important;
        box-sizing: border-box !important;
    }
    div[data-testid="stDownloadButton"] > button:hover,
    .stDownloadButton > button:hover {
        background-color: #475569 !important;
        border-color: #38bdf8 !important;
        color: #ffffff !important;
        box-shadow: 0 0 10px rgba(56, 189, 248, 0.25) !important;
    }
    div[data-testid="stDownloadButton"] > button:active,
    .stDownloadButton > button:active {
        background-color: #1e293b !important;
        border-color: #0284c7 !important;
    }
    div[data-testid="stDownloadButton"] > button p,
    div[data-testid="stDownloadButton"] > button span,
    .stDownloadButton > button p,
    .stDownloadButton > button span {
        font-size: 0.875rem !important;
        font-weight: 500 !important;
        color: inherit !important;
        line-height: inherit !important;
        margin: 0 !important;
        padding: 0 !important;
    }

    /* 사이드바 스타일링 */
    section[data-testid="stSidebar"], [data-testid="stSidebar"] {
        background-color: #1e293b !important;
        border-right: 1px solid #334155 !important;
    }
    section[data-testid="stSidebar"] h1,
    section[data-testid="stSidebar"] h2,
    section[data-testid="stSidebar"] h3 {
        color: #f8fafc !important;
        -webkit-text-fill-color: #f8fafc !important;
    }

    /* =========================================================
       사이드바 접기(<<) 및 펼치기(>>) 버튼 항상 표시 및 시인성/대비 강화
       ========================================================= */
    /* 1. 사이드바가 열려 있을 때 접기 버튼 (<<) 상시 표시 */
    [data-testid="stSidebarCollapseButton"] {
        visibility: visible !important;
        opacity: 1 !important;
        display: inline-flex !important;
    }
    
    [data-testid="stSidebarCollapseButton"] button {
        visibility: visible !important;
        opacity: 1 !important;
        background-color: #1e293b !important;       /* 진한 네이비 배경 */
        border: 1.5px solid #38bdf8 !important;     /* 선명한 스카이블루 테두리로 상자 명확화 */
        border-radius: 8px !important;
        width: 38px !important;
        height: 38px !important;
        display: flex !important;
        align-items: center !important;
        justify-content: center !important;
        box-shadow: 0 2px 8px rgba(0, 0, 0, 0.4), 0 0 6px rgba(56, 189, 248, 0.2) !important;
        transition: all 0.2s ease !important;
    }
    
    /* 상자 내부의 << 아이콘(Material Icon span/svg/문자)을 순백색으로 강제하여 상자와 극명한 대비 구현 */
    [data-testid="stSidebarCollapseButton"] button *,
    [data-testid="stSidebarCollapseButton"] span,
    [data-testid="stSidebarCollapseButton"] [data-testid="stIconMaterial"],
    [data-testid="stSidebarCollapseButton"] svg {
        color: #ffffff !important;
        fill: #ffffff !important;
        opacity: 1 !important;
        visibility: visible !important;
        font-size: 1.35rem !important;
        font-weight: 700 !important;
    }
    
    /* 호버(PC) 및 터치 시 반전 효과 */
    [data-testid="stSidebarCollapseButton"] button:hover {
        background-color: #38bdf8 !important;
        border-color: #38bdf8 !important;
    }
    [data-testid="stSidebarCollapseButton"] button:hover * {
        color: #0f172a !important;
        fill: #0f172a !important;
    }

    /* 2. 사이드바 헤더 영역 패딩 및 정렬 보정 */
    [data-testid="stSidebarHeader"] {
        padding-top: 0.5rem !important;
        padding-bottom: 0.5rem !important;
    }

    /* 3. 사이드바가 닫혔을 때 다시 여는 버튼 (>>) 시인성 강화 */
    [data-testid="stSidebarCollapsedControl"] {
        visibility: visible !important;
        opacity: 1 !important;
    }
    
    [data-testid="stSidebarCollapsedControl"] button {
        background-color: #1e293b !important;
        border: 1.5px solid #38bdf8 !important;
        border-radius: 8px !important;
        box-shadow: 0 2px 8px rgba(0, 0, 0, 0.4), 0 0 6px rgba(56, 189, 248, 0.2) !important;
    }
    
    [data-testid="stSidebarCollapsedControl"] button *,
    [data-testid="stSidebarCollapsedControl"] span,
    [data-testid="stSidebarCollapsedControl"] [data-testid="stIconMaterial"],
    [data-testid="stSidebarCollapsedControl"] svg {
        color: #38bdf8 !important;
        fill: #38bdf8 !important;
        opacity: 1 !important;
        visibility: visible !important;
        font-size: 1.35rem !important;
    }

    /* Primary Button Styling (39 DividendStock 테마 통일) */
    .stButton button[kind="primary"],
    .stButton > button[kind="primary"],
    section[data-testid="stSidebar"] button[kind="primary"] {
        background-color: #2563eb !important;
        color: #ffffff !important;
        border: none !important;
        font-weight: 600 !important;
        border-radius: 6px !important;
        transition: all 0.2s ease !important;
    }
    .stButton button[kind="primary"]:hover,
    .stButton > button[kind="primary"]:hover,
    section[data-testid="stSidebar"] button[kind="primary"]:hover {
        background-color: #1d4ed8 !important;
        box-shadow: 0 0 10px rgba(37, 99, 235, 0.4) !important;
    }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-title">David Ryan "Just Draw the Line" 스크리너</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-title">한국 및 미국 주식시장의 종목 중 추세 돌파 및 거래량 동반 종목 발굴 프로그램</div>', unsafe_allow_html=True)

# 기법 소개
with st.expander("ℹ️ 데이비드 라이언의 'Just Draw the Line' 투자 기법이란?"):
    st.markdown("""
    **데이비드 라이언(David Ryan)**은 윌리엄 오닐의 제자이자, 미국 투자 챔피언십 3년 연속 우승에 빛나는 전설적인 투자자입니다.
    
    그는 차트를 지나치게 복잡한 지표(RSI, MACD 등)로 어지럽히지 않고, **오직 주가와 거래량**에 집중하여 선을 그릴 것을 강조했습니다.
    
    ### 📌 핵심 스크리닝 요건
    1. **상승 추세 (Stage 2 Uptrend) 확인**:
       * 주가가 50일, 150일, 200일 이동평균선 위에 위치.
       * 이평선 정배열 (50MA > 150MA > 200MA).
       * 200일 이평선이 최소 1개월 동안 상승 흐름 유지.
       * 주가가 52주 신저가 대비 최소 25% 이상 높고, 52주 신고가 대비 25% 이내에 위치 (박스권 상단 대기).
    2. **하향 추세선 돌파 (Just Draw the Line)**:
       * 최근 하락 조정 기간 동안 고점들을 연결한 상단 저항선(Downtrend line)을 도출.
       * 당일(혹은 직전 영업일) 주가가 이 추세선을 **상향 돌파(Breakout)**하여 마감.
    3. **거래량 확인 (Volume Confirmation)**:
       * 돌파 시점의 거래량이 **최근 20일 평균 거래량 대비 최소 1.5배(150%) 이상** 급증하여 기관의 매수세 확인.
    """)

# 세션 상태 초기화 (스크리닝 결과 보존용)
if 'screened_df' not in st.session_state:
    st.session_state.screened_df = None
if 'last_run_time' not in st.session_state:
    st.session_state.last_run_time = None
if 'market_type_used' not in st.session_state:
    st.session_state.market_type_used = None

# 사이드바 설정 영역
with st.sidebar:
    st.markdown(
        """
        <div style='padding: 2px 0 12px 0;'>
            <div style='font-size: 1.25rem; font-weight: 700; color: #f8fafc; letter-spacing: -0.01em; display: flex; align-items: center; gap: 8px;'>
                <span>⚙️</span> 스크리닝 조건 설정
            </div>
            <div style='font-size: 0.82rem; color: #94a3b8; margin-top: 4px; line-height: 1.4;'>
                데이비드 라이언 추세선 돌파 분석 조건을 설정합니다.
            </div>
        </div>
        <hr style='border: 0; height: 1px; background-color: #334155; margin: 10px 0 16px 0;'>
        """,
        unsafe_allow_html=True
    )

    market_choice = st.selectbox(
        "🏛️ 시장 선택",
        ["KOSPI", "KOSDAQ", "S&P 500", "NASDAQ 100"],
        index=0
    )

    # 한국 시장(KOSPI, KOSDAQ)일 경우 시가총액 기반 대상 범위(Scope) 및 최소 시총 옵션 제공
    is_korean_market = market_choice in ["KOSPI", "KOSDAQ"]
    scope_code = "top500"
    min_marcap_val = 0
    if is_korean_market:
        scope_options = {
            "시총 상위 300 (쾌속 모드 ~15초)": "top300",
            "시총 상위 500 (권장 모드 ~25초)": "top500",
            "시총 상위 1,000 (심층 모드 ~50초)": "top1000",
            "시장 전체 종목 (전체 모드)": "all"
        }
        scope_choice_label = st.selectbox(
            "🎯 대상 범위 (Scope)",
            options=list(scope_options.keys()),
            index=1,
            help="시가총액 상위 종목 위주로 분석하여 스크리닝 속도를 최적화합니다."
        )
        scope_code = scope_options[scope_choice_label]

        marcap_options = {
            "제한 없음 (전체)": 0,
            "1,000억원 이상": 1000,
            "3,000억원 이상 [추천]": 3000,
            "5,000억원 이상": 5000,
            "1조원 이상": 10000
        }
        marcap_label = st.selectbox(
            "💰 최소 시가총액",
            options=list(marcap_options.keys()),
            index=0,
            help="설정한 시가총액 이상의 종목만 스크리닝합니다."
        )
        min_marcap_val = marcap_options[marcap_label]

    st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True)
    start_screening = st.button("🔍 스크리닝 시작", type="primary", use_container_width=True)

    st.markdown("<hr style='border: 0; height: 1px; background-color: #334155; margin: 16px 0;'>", unsafe_allow_html=True)
    st.markdown("<div style='font-size: 0.95rem; font-weight: 700; color: #e2e8f0; margin-bottom: 6px;'>🎯 스크리닝 필터 설정</div>", unsafe_allow_html=True)

    lookback_period = st.slider(
        "추세선 분석 기간 (영업일)",
        min_value=20,
        max_value=90,
        value=40,
        step=5,
        help="최근 고점을 연결하여 추세선을 그릴 분석 윈도우 기간입니다."
    )

    vol_ratio_thresh = st.slider(
        "최소 돌파 거래량 배수",
        min_value=1.0,
        max_value=3.0,
        value=1.5,
        step=0.1,
        help="돌파 당일 거래량이 직전 20일 평균 거래량 대비 몇 배 이상이어야 하는지 결정합니다. (예: 1.5 = 150%)"
    )

    apply_trend_template = st.toggle(
        "장기 상승 추세 조건(Trend Template) 필터",
        value=True,
        help="미너비니의 상승 2단계 정배열 조건을 활성화합니다. 활성화하면 매우 엄격한 상승 추세 종목만 발굴됩니다."
    )

    breakout_window = st.slider(
        "최근 돌파 허용 기간 (영업일)",
        min_value=1,
        max_value=10,
        value=3,
        step=1,
        help="최근 N영업일 이내에 최초 돌파가 일어난 후 추세선 위를 지키고 있는 종목을 허용합니다."
    )

if start_screening:
    market_map = {
        "KOSPI": "KOSPI",
        "KOSDAQ": "KOSDAQ",
        "S&P 500": "S&P 500",
        "NASDAQ 100": "NASDAQ 100",
        "NASDAQ": "NASDAQ 100",
        "코스피 (KOSPI)": "KOSPI",
        "코스닥 (KOSDAQ)": "KOSDAQ",
        "전체 시장 (KOSPI + KOSDAQ)": "ALL",
        "미국 S&P 500 (US)": "S&P 500",
        "미국 NASDAQ 100 (US)": "NASDAQ 100"
    }
    selected_market = market_map.get(market_choice, "KOSPI")
    
    with st.spinner("상장 종목 유니버스를 로드하는 중..."):
        try:
            tickers_df = get_krx_tickers(selected_market, scope=scope_code, min_marcap_eok=min_marcap_val)
            total_count = len(tickers_df)
            st.info(f"수집 대상 유니버스: 총 {total_count}개 종목 (노이즈 필터링 완료)")
        except Exception as e:
            st.error(f"종목 목록 수집 실패: {e}")
            tickers_df = pd.DataFrame()
            
    if not tickers_df.empty:
        progress_bar = st.progress(0.0)
        status_text = st.empty()
        
        def update_progress(current, total, name):
            ratio = min(1.0, current / total) if total > 0 else 0.0
            progress_bar.progress(ratio)
            status_text.markdown(f"⏳ **데이터 다운로드 및 추세선 돌파 분석 중...** ({current}/{total}) `{name}`")

        start_time = time.time()
        try:
            with st.spinner("초고속 멀티스레딩 데이터 수집 및 선형계획법(linprog) 병렬 연산 중..."):
                df_screened = run_screening_task(
                    tickers_df=tickers_df,
                    lookback_period=lookback_period,
                    vol_ratio_thresh=vol_ratio_thresh,
                    apply_trend_template=apply_trend_template,
                    breakout_window=breakout_window,
                    max_workers=24,
                    progress_callback=update_progress
                )
                elapsed = time.time() - start_time
                progress_bar.progress(1.0)
                if df_screened is not None and not df_screened.empty:
                    status_text.success(f"✅ 스크리닝 완료! ({len(df_screened)}개 종목 발굴, 소요 시간: {elapsed:.1f}초)")
                else:
                    status_text.warning(f"⚠️ 조건에 부합하는 종목이 없습니다. (소요 시간: {elapsed:.1f}초)")
                time.sleep(0.8)
                progress_bar.empty()
                status_text.empty()
        except Exception as e:
            st.error(f"스크리닝 작업 중 오류 발생: {e}")
            df_screened = pd.DataFrame()
            progress_bar.empty()
            status_text.empty()

        if df_screened is not None and not df_screened.empty:
            df_final = df_screened.copy()
            # 출력용 한글 칼럼명 매핑
            df_final_display = df_final.rename(columns={
                'ticker': '티커',
                'name': '종목명',
                'price': '현재가',
                'ma50': '50일 MA',
                'ma150': '150일 MA',
                'ma200': '200일 MA',
                'high_52w': '52주 최고가',
                'low_52w': '52주 최저가',
                'vol_ratio': '거래량 비율',
                'breakout_date': '돌파 감지일'
            })
            # 불필요한 칼럼 제거
            df_final_display = df_final_display.drop(columns=['trend_slope', 'trend_intercept'], errors='ignore')
            
            st.session_state.screened_df = df_final_display
            st.session_state.raw_screened_df = df_final # 원본 저장
        else:
            st.session_state.screened_df = pd.DataFrame()
            st.session_state.raw_screened_df = pd.DataFrame()
            
        st.session_state.last_run_time = datetime.datetime.now(KST).strftime('%Y-%m-%d %H:%M:%S')
        st.session_state.market_type_used = market_choice

# 결과 디스플레이
if st.session_state.screened_df is not None:
    st.success(f"🔍 스크리닝 완료! (실행 시각: {st.session_state.last_run_time} | 대상: {st.session_state.market_type_used})")
    
    if st.session_state.screened_df.empty:
        st.warning("조건에 부합하는 종목이 발견되지 않았습니다. 분석 기간을 늘리거나 거래량 배수를 낮춰 보세요.")
    else:
        # --- 엑셀 저장용 데이터 사전 가공 및 생성 ---
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # 엑셀 내보내기용 별도 데이터프레임 가공 (raw_screened_df 기반)
            df_excel = st.session_state.raw_screened_df.copy()
            df_excel = df_excel.rename(columns={
                'ticker': '티커',
                'name': '종목명',
                'price': '현재가',
                'ma50': '50일 MA',
                'ma150': '150일 MA',
                'ma200': '200일 MA',
                'high_52w': '52주 최고가',
                'low_52w': '52주 최저가',
                'vol_ratio': '거래량 비율',
                'breakout_date': '돌파 감지일'
            })
            
            # 한국/미국 주식 구분에 따른 소수점 라운딩 및 형변환
            def format_excel_data(row):
                ticker = row['티커']
                is_kr = ticker.endswith('.KS') or ticker.endswith('.KQ')
                
                # 가격 관련 필드들
                price_cols = ['현재가', '50일 MA', '150일 MA', '200일 MA', '52주 최고가', '52주 최저가']
                for col in price_cols:
                    if is_kr:
                        # 한국 주식은 소수점 반올림 후 정수로 변환
                        row[col] = int(round(row[col]))
                    else:
                        # 미국 주식은 소수점 둘째 자리까지 반올림
                        row[col] = round(row[col], 2)
                        
                # 거래량 비율은 공통 소수점 둘째 자리 반올림
                row['거래량 비율'] = round(row['거래량 비율'], 2)
                return row
                
            df_excel = df_excel.apply(format_excel_data, axis=1)
            # 불필요한 분석용 내부 칼럼 제외
            df_excel = df_excel.drop(columns=['trend_slope', 'trend_intercept'], errors='ignore')
            
            # 가공된 데이터프레임을 엑셀에 쓰기
            df_excel.to_excel(writer, index=False, sheet_name='Just Draw the Line')
            
            # openpyxl 객체 제어로 엑셀 서식화
            worksheet = writer.sheets['Just Draw the Line']
            
            # 1. 1행 헤더에 필터/정렬 토글(AutoFilter) 적용
            from openpyxl.utils import get_column_letter
            max_col = worksheet.max_column
            max_row = worksheet.max_row
            if max_row > 0:
                worksheet.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
                
            # 2. 열 너비 자동 맞춤 (Auto-fit) 적용
            for col in worksheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val = str(cell.value or '')
                    # 한글 문자 폭(전각 문자) 보정을 고려한 글자수 산출
                    length = sum(2 if ord(char) > 128 else 1 for char in val)
                    if length > max_len:
                        max_len = length
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
                
            # 3. 특정 열 서식 및 정렬 설정 추가 (미국 주식 가격 필드 및 거래량 비율 '0.00' 서식 적용, J열 가운데 정렬)
            from openpyxl.styles import Alignment
            for row_idx in range(2, max_row + 1):
                # A열(티커 - 1번째 열)의 값 분석
                ticker_val = str(worksheet.cell(row=row_idx, column=1).value or '')
                is_kr = ticker_val.endswith('.KS') or ticker_val.endswith('.KQ')
                
                # 미국 주식(달러화 자산)일 경우 가격 열(C~H열, 즉 3~8번째 열)에 소수점 2자리 '0.00' 서식 적용
                if not is_kr:
                    for col_idx in range(3, 9): # 3열(현재가) ~ 8열(52주 최저가)
                        worksheet.cell(row=row_idx, column=col_idx).number_format = '0.00'
                
                # I열 (거래량 비율 - 9번째 열) -> 공통 소수점 2자리 '0.00' 서식 적용
                cell_i = worksheet.cell(row=row_idx, column=9)
                cell_i.number_format = '0.00'
                
                # J열 (돌파 감지일 - 10번째 열) -> 가운데 정렬 적용
                cell_j = worksheet.cell(row=row_idx, column=10)
                cell_j.alignment = Alignment(horizontal='center')
                
        excel_data = output.getvalue()
        
        # 파일명 동적 생성 (JustDrawLine-[MarketCode]-YYYY-MM-DD.xlsx)
        market_code_map = {
            # 현재 사이드바 UI 선택값
            "KOSPI": "KS",
            "KOSDAQ": "KQ",
            "S&P 500": "SP",
            "NASDAQ 100": "NQ",
            "NASDAQ": "NQ",
            # 레거시 및 호환용 명칭
            "코스피 (KOSPI)": "KS",
            "코스닥 (KOSDAQ)": "KQ",
            "전체 시장 (KOSPI + KOSDAQ)": "KS&KQ",
            "미국 S&P 500 (US)": "SP",
            "미국 NASDAQ 100 (US)": "NQ",
            "ALL": "KS&KQ"
        }
        raw_market = str(st.session_state.market_type_used).strip() if st.session_state.market_type_used else ""
        market_code = market_code_map.get(raw_market)
        if not market_code:
            key_upper = raw_market.upper()
            if "KOSPI" in key_upper or "코스피" in raw_market:
                market_code = "KS"
            elif "KOSDAQ" in key_upper or "코스닥" in raw_market:
                market_code = "KQ"
            elif "S&P" in key_upper or "500" in raw_market:
                market_code = "SP"
            elif "NASDAQ" in key_upper or "100" in raw_market:
                market_code = "NQ"
            else:
                market_code = "ALL"
        today_str = datetime.datetime.now(KST).strftime('%Y-%m-%d')
        excel_filename = f"JustDrawLine-{market_code}-{today_str}.xlsx"

        # 타이틀과 엑셀 다운로드 버튼을 같은 라인에 배치 (다운로드 버튼은 오른쪽 끝에 정렬)
        col_title, col_btn = st.columns([8, 2], vertical_alignment="bottom")
        with col_title:
            st.markdown(
                f"<div style='font-size: 1.20rem; font-weight: 700; color: #8AB4F8; margin: 10px 0 6px 0; display: flex; align-items: center; gap: 8px;'>"
                f"<span>📋</span> 스크리닝 결과 (총 {len(st.session_state.screened_df)}개 종목)"
                f"</div>",
                unsafe_allow_html=True
            )
        with col_btn:
            st.download_button(
                label="📥 엑셀 파일 다운로드",
                data=excel_data,
                file_name=excel_filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        
        # 소수점 포맷팅
        df_format = st.session_state.screened_df.copy()
        
        df_format['현재가'] = df_format.apply(lambda r: fmt_curr(r['현재가'], r['티커']), axis=1)
        df_format['50일 MA'] = df_format.apply(lambda r: fmt_curr(r['50일 MA'], r['티커']), axis=1)
        df_format['150일 MA'] = df_format.apply(lambda r: fmt_curr(r['150일 MA'], r['티커']), axis=1)
        df_format['200일 MA'] = df_format.apply(lambda r: fmt_curr(r['200일 MA'], r['티커']), axis=1)
        df_format['52주 최고가'] = df_format.apply(lambda r: fmt_curr(r['52주 최고가'], r['티커']), axis=1)
        df_format['52주 최저가'] = df_format.apply(lambda r: fmt_curr(r['52주 최저가'], r['티커']), axis=1)
        df_format['거래량 비율'] = df_format['거래량 비율'].map('{:.2f}배'.format)

        st.dataframe(df_format, use_container_width=True)
        
        # --- 개별 종목 차트 시각화 영역 ---
        st.markdown("---")
        st.markdown(
            "<div style='font-size: 1.20rem; font-weight: 700; color: #8AB4F8; margin: 20px 0 10px 0; display: flex; align-items: center; gap: 8px;'>"
            "<span>📈</span> 종목별 추세선 분석 차트"
            "</div>",
            unsafe_allow_html=True
        )
        
        # 사용자가 차트로 확인해볼 종목 선택
        selected_stock_name = st.selectbox(
            "시각화할 종목을 선택하세요",
            options=st.session_state.screened_df['종목명'].tolist()
        )
        
        if selected_stock_name:
            # 선택된 종목의 원본 행 데이터 찾기
            row = st.session_state.raw_screened_df[st.session_state.raw_screened_df['name'] == selected_stock_name].iloc[0]
            ticker = row['ticker']
            
            with st.spinner(f"{selected_stock_name} ({ticker}) 주가 데이터 가져오는 중..."):
                # 차트 작성을 위해 2년치 데이터 수집 (여유있게 MA를 그리기 위함)
                is_kr = ticker.endswith('.KS') or ticker.endswith('.KQ')
                if is_kr:
                    code = ticker.split('.')[0]
                    start_date = (datetime.datetime.now() - datetime.timedelta(days=730)).strftime('%Y-%m-%d')
                    df_chart = fdr.DataReader(code, start_date)
                else:
                    df_chart = yf.download(ticker, period="2y", progress=False)
                    if isinstance(df_chart.columns, pd.MultiIndex):
                        df_chart.columns = df_chart.columns.droplevel(1)

                df_chart = df_chart.dropna(subset=['Close', 'High', 'Low', 'Volume'])
                if df_chart.index.tz is not None:
                    df_chart.index = df_chart.index.tz_localize(None)

                
            if not df_chart.empty:
                # 차트용 데이터 계산
                close_prices = df_chart['Close']
                high_prices = df_chart['High']
                low_prices = df_chart['Low']
                volume_prices = df_chart['Volume']
                
                ma50 = close_prices.rolling(window=50).mean()
                ma150 = close_prices.rolling(window=150).mean()
                ma200 = close_prices.rolling(window=200).mean()
                
                # 최근 lookback_period개의 고가로 다시 추세선 피팅
                recent_highs = high_prices.iloc[-lookback_period:].values
                trendline, a, b = fit_upper_trendline(recent_highs)
                
                # 전체 인덱스 중 최근 lookback_period에 해당하는 날짜 목록
                recent_dates = df_chart.index[-lookback_period:]
                
                # Plotly 서브플롯 생성 (주가 캔들스틱 + 거래량 바)
                fig = make_subplots(
                    rows=2, cols=1, 
                    shared_xaxes=True, 
                    vertical_spacing=0.08,
                    row_heights=[0.7, 0.3]
                )
                
                # 1. 캔들스틱 차트 추가
                fig.add_trace(
                    go.Candlestick(
                        x=df_chart.index,
                        open=df_chart['Open'],
                        high=df_chart['High'],
                        low=df_chart['Low'],
                        close=df_chart['Close'],
                        name="주가",
                        increasing_line_color='#EA4335', # 한국 스타일 빨간색 상승
                        decreasing_line_color='#4285F4'  # 한국 스타일 파란색 하락
                    ),
                    row=1, col=1
                )
                
                # 2. 이동평균선 추가
                fig.add_trace(
                    go.Scatter(x=df_chart.index, y=ma50, line=dict(color='#FBBC05', width=1.5), name="50일 MA"),
                    row=1, col=1
                )
                fig.add_trace(
                    go.Scatter(x=df_chart.index, y=ma150, line=dict(color='#34A853', width=1.5), name="150일 MA"),
                    row=1, col=1
                )
                fig.add_trace(
                    go.Scatter(x=df_chart.index, y=ma200, line=dict(color='#EA4335', width=2), name="200일 MA"),
                    row=1, col=1
                )
                
                # 3. "Just Draw the Line" 하향 추세선 오버레이
                if trendline is not None:
                    fig.add_trace(
                        go.Scatter(
                            x=recent_dates,
                            y=trendline,
                            line=dict(color='#FFFFFF', width=3, dash='dash'),
                            name="하향 추세선 (Downtrend Resistance Line)"
                        ),
                        row=1, col=1
                    )
                    
                    # 돌파 지점 하이라이트 (돌파 감지일 기준 동적 위치 탐색)
                    breakout_date_dt = pd.to_datetime(row['breakout_date'])
                    try:
                        breakout_idx = df_chart.index.get_loc(breakout_date_dt)
                        breakout_price = df_chart['Close'].iloc[breakout_idx]
                    except Exception:
                        # 매칭 실패 시 차트 마지막 일자로 대체
                        breakout_idx = -1
                        breakout_date_dt = df_chart.index[-1]
                        breakout_price = df_chart['Close'].iloc[-1]

                    
                    fig.add_annotation(
                        x=breakout_date_dt,
                        y=breakout_price,
                        text="★ 돌파 (Breakout)",
                        showarrow=True,
                        arrowhead=2,
                        arrowsize=1,
                        arrowwidth=2,
                        arrowcolor="#EA4335",
                        ax=0,
                        ay=-40,
                        font=dict(color="#EA4335", size=12, family="Malgun Gothic"),
                        row=1, col=1
                    )
                
                # 4. 거래량 차트 추가
                # 상승/하락일에 따른 거래량 색상 구분
                colors = ['#EA4335' if df_chart['Close'].iloc[i] >= df_chart['Open'].iloc[i] else '#4285F4' for i in range(len(df_chart))]
                fig.add_trace(
                    go.Bar(
                        x=df_chart.index,
                        y=df_chart['Volume'],
                        marker_color=colors,
                        name="거래량"
                    ),
                    row=2, col=1
                )
                
                # 거래량 20일 이동평균 추가
                vol_ma20 = df_chart['Volume'].rolling(window=20).mean()
                fig.add_trace(
                    go.Scatter(
                        x=df_chart.index,
                        y=vol_ma20,
                        line=dict(color='#5F6368', width=1.5),
                        name="20일 거래량 MA"
                    ),
                    row=2, col=1
                )
                
                # 화폐 단위 동적 결정
                is_us_stock = not (ticker.endswith('.KS') or ticker.endswith('.KQ'))
                currency_symbol = '$' if is_us_stock else '원'
                tick_format = ',.2f' if is_us_stock else ',.0f'
                
                # 레이아웃 정밀화 (고대비 Tailwind Slate 표준 테마)
                fig.update_layout(
                    template="plotly_dark",
                    paper_bgcolor=STANDARD_CHART_THEME['paper_bgcolor'],
                    plot_bgcolor=STANDARD_CHART_THEME['plot_bgcolor'],
                    title=dict(
                        text=f"<b>📈 {selected_stock_name} ({ticker}) 'Just Draw the Line' 분석 차트</b>",
                        font=dict(color="#F8FAFC", size=16)
                    ),
                    yaxis_title=f"주가 ({currency_symbol})",
                    yaxis2_title="거래량 (주)",
                    xaxis_rangeslider_visible=False,
                    height=700,
                    margin=dict(l=50, r=50, t=80, b=50),
                    showlegend=True,
                    legend=dict(
                        orientation="h",
                        y=1.08,
                        xanchor="right",
                        x=1,
                        bgcolor="rgba(30, 41, 59, 0.85)",
                        bordercolor="#334155",
                        borderwidth=1,
                        font=dict(color="#F8FAFC", size=11)
                    ),
                    hovermode="x unified"
                )
                
                fig.update_xaxes(
                    tickformat="%Y-%m-%d",
                    hoverformat="%Y-%m-%d",
                    gridcolor="#334155", 
                    linecolor="#475569", 
                    tickfont=dict(color="#cbd5e1")
                )
                fig.update_yaxes(tickformat=tick_format, row=1, col=1, gridcolor="#334155", linecolor="#475569", tickfont=dict(color="#cbd5e1"))
                fig.update_yaxes(tickformat=",.0f", row=2, col=1, gridcolor="#334155", linecolor="#475569", tickfont=dict(color="#cbd5e1"))
                
                # 주말 및 공휴일 공백 제거 (5일 주기 끊김 및 0값 방지)
                dt_all = pd.date_range(start=df_chart.index[0], end=df_chart.index[-1], freq='B')
                existing_dates = set(pd.to_datetime(df_chart.index).normalize())
                holidays = [d.strftime('%Y-%m-%d') for d in dt_all if d.normalize() not in existing_dates]
                rbreaks = [dict(bounds=["sat", "mon"])]
                if holidays:
                    rbreaks.append(dict(values=holidays))
                fig.update_xaxes(rangebreaks=rbreaks)
                
                st.plotly_chart(fig, use_container_width=True)
                
                # 추가 설명 카드
                st.markdown(f"""
                <div class="metric-card">
                    <div style="font-size: 1.00rem; font-weight: 600; color: #CBD5E1; margin-bottom: 10px; display: flex; align-items: center; gap: 6px;">
                        <span>💡</span> {selected_stock_name} 상세 분석 정보
                    </div>
                    <ul>
                        <li><b>돌파 발생일:</b> {row['breakout_date']}</li>
                        <li><b>돌파 시점 거래량 폭증 비율:</b> <span style="color:#EA4335; font-weight:bold;">{row['vol_ratio']:.2f}배</span> (이전 20일 평균 거래량 대비)</li>
                        <li><b>52주 최고가 대비 가격:</b> {row['price'] / row['high_52w'] * 100:.1f}% 수준 (최고가: {fmt_curr(row['high_52w'], ticker)})</li>
                        <li><b>52주 최저가 대비 가격:</b> +{ (row['price'] / row['low_52w'] - 1) * 100:.1f}% 상승 상태 (최저가: {fmt_curr(row['low_52w'], ticker)})</li>
                    </ul>
                </div>
                """, unsafe_allow_html=True)

else:
    # 프로그램 최초 진입 시 메인 화면
    st.info("👈 왼쪽 사이드바에서 대상 시장 및 파라미터를 설정한 후 '스크리닝 시작' 버튼을 눌러주세요.")

st.markdown("---")
st.markdown("<div style='text-align: center; color: #64748b; font-size: 0.8rem; margin-top: 8px; margin-bottom: 24px; line-height: 1.6;'>⚠️ 본 서비스에서 제공하는 모든 정보는 투자 참고용이며, 투자의 최종 결정과 책임은 투자자 본인에게 있습니다.</div>", unsafe_allow_html=True)
